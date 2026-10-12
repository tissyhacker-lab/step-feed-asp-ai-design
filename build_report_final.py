from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_report import (
    GRAY,
    NAVY,
    OUT,
    add_caption,
    add_equation as _add_equation,
    add_heading,
    add_page_number,
    add_table as _add_table,
)


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "HU_TONGYU_MA_DINGNA_Final.docx"
GITHUB_URL = "https://github.com/tissyhacker-lab/step-feed-asp-ai-design"
EQUATION_NUMBER = 0


def blank_line(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("").font.size = Pt(9)


def add_table(doc: Document, *args, **kwargs):
    blank_line(doc)
    table = _add_table(doc, *args, **kwargs)
    # Keep compact tables together when enough room exists on the next page.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    blank_line(doc)
    return table


def add_equation(doc: Document, equation: str, note: str = "") -> None:
    global EQUATION_NUMBER
    EQUATION_NUMBER += 1
    blank_line(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    r = p.add_run(equation)
    r.font.name = "Cambria Math"
    r.font.size = Pt(10.2)
    if note:
        rn = p.add_run(f"  {note}")
        rn.font.name = "Arial"
        rn.font.size = Pt(8.2)
    p.add_run("\t")
    n = p.add_run(f"({EQUATION_NUMBER})")
    n.font.name = "Arial"
    n.font.size = Pt(9)
    blank_line(doc)


def body(doc: Document, text: str, lead: str | None = None, size: float = 9.0, after: float = 2.0) -> None:
    """Assignment body text: explicitly double-spaced."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(after)
    if lead and text.startswith(lead):
        r = p.add_run(lead)
        r.bold = True
        p.add_run(text[len(lead) :])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)


def compact_note(doc: Document, text: str, size: float = 7.7) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(GRAY)


def page(doc: Document) -> None:
    # Sections now flow continuously; Word paginates by content instead of Step.
    return


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.15)
    sec.bottom_margin = Cm(1.15)
    sec.left_margin = Cm(1.35)
    sec.right_margin = Cm(1.35)
    add_page_number(sec)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(9.0)
    for level, size in ((1, 14), (2, 10.5)):
        st = styles[f"Heading {level}"]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY)

    # Page 1: identity, scope and auditable design basis.
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("AI-Assisted Feed-Split Design for a Full-Scale Step-Feed ASP")
    r.font.name = "Arial"
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    authors = doc.add_paragraph()
    authors.paragraph_format.space_after = Pt(5)
    ar = authors.add_run("HU TONGYU, A0332797N  |  MA DINGNA, A0319575R")
    ar.font.name = "Arial"
    ar.font.size = Pt(10.5)
    ar.font.bold = True

    add_heading(doc, "1  Design question and case basis")
    body(
        doc,
        "This study asks whether AI can improve the normal-operation influent split among four equal-volume zones of a real step-feed activated-sludge design. The 2021 EPA San Antonio de los Buenos feasibility case supplies the hydraulic, organic and total-aeration basis; the EPA 25/25/25/25 split is withheld until the engineering reveal. AI generates alternatives, a mechanistic model labels them, a machine-learning surrogate screens them, and engineering rules determine what can be recommended.",
    )
    add_table(
        doc,
        ["Design basis", "Value", "Audit status"],
        [
            ["Average / peak flow", "5.0 / 12.5 MGD", "EPA Tables 2-1, 2-7"],
            ["Secondary-influent BOD5", "11,700 lb/d = 280.58 mg/L", "8.34 conversion included"],
            ["Aeration volume / MLSS", "2.0 MG / 3,000 mg/L", "Two 1.0-MG reactors"],
            ["RAS ratio", "R = 2.1/5.0 = 0.42", "EPA design flow"],
            ["O2 required / supplied", "15,400 / 15,600 lb/d", "Total plant values only"],
            ["Decision variables", "f1...f4; sum(fi)=1", "0.05 <= fi <= 0.60"],
        ],
        [6.0, 6.2, 5.9],
        8.0,
    )
    add_heading(doc, "Independent design checks", level=2)
    add_equation(doc, "theta = 24V/Q = 24(2.0)/5.0 = 9.6 h")
    add_equation(doc, "L_v = 8.34QS_0/(133.681V) = 43.76 lb BOD5 d^-1 (10^3 ft^3)^-1")
    add_equation(doc, "F/M = QS_0/(VX) = 0.234 d^-1")
    body(
        doc,
        "The corrected loading equation contains 8.34 because Q is in MGD and S0 is in mg/L; 133.681 converts MG to 10^3 ft3. The checks reproduce the reported design scale and establish transparent guardrails. They do not establish current as-built performance.",
        size=8.7,
    )
    add_heading(doc, "AI role and claimed contribution", level=2)
    body(
        doc,
        "AI creativity is expressed through divergent option generation and discovery of a local-aeration-capacity frontier, not through inventing biological laws. Its productivity is the rapid evaluation of thousands of four-way splits and automatic sensitivity analysis. Effectiveness is judged by reproducibility, mechanistic recomputation and the ability to expose an engineering decision that a single hand calculation would hide.",
        size=8.7,
    )

    # Page 2: model, oxygen bookkeeping and ML evidence boundary.
    page(doc)
    add_heading(doc, "2  Model, sampling and surrogate evidence")
    body(
        doc,
        "A four-CSTR steady-state model represents the aggregate four zones. Each aggregate zone uses V_i = 0.50 MG, i = 1, 2, 3, 4. Soluble BOD in recycled activated sludge is approximated by the final-zone concentration. The one-point rate coefficient is calibrated from the EPA aggregate influent load, volume and 11 mg/L modelled cBOD5 value. This makes the model suitable for relative split screening under fixed assumptions, not for independent prediction of plant effluent.",
    )
    add_table(
        doc,
        ["Element", "Equation", "Meaning"],
        [
            ["Zone 1 boundary", "0=RQ S4+f1QS0-Q(R+f1)S1-kappa V1S1", "RAS soluble BOD enters as RQ S4"],
            ["Zones 2-4", "0=q(i-1)S(i-1)+fiQS0-qiSi-kappa ViSi", "i=2,3,4; Vi=0.50 MG"],
            ["Cumulative flow", "qi=Q(R+sum[j<=i]fj)", "Step feed plus RAS"],
            ["Calibration", "kappa=Q(S0-Se)/(VSe)=61.27 d^-1", "One-point aggregate calibration"],
            ["BOD removed", "DeltaLi=8.34 kappa Vi Si", "lb BOD5/d"],
            ["Zone oxygen", "Oi=1.1 DeltaLi+Obackground/4", "Screening oxygen allocation"],
            ["Objective", "0.50 Se/15+0.30P+0.15CV+0.05 Ototal/15600", "P and CV defined below"],
        ],
        [3.1, 9.1, 5.9],
        7.7,
    )
    add_equation(doc, "O_bar=(sum O_i)/4;  P=max(O_i)/O_bar;  CV=sqrt[sum(O_i-O_bar)^2/4]/O_bar")
    add_heading(doc, "Reproducible background oxygen term", level=2)
    add_equation(doc, "O_background = 15,400 - 1.1 sum(DeltaL_i,equal) = 3,034.57 lb O2/d")
    body(
        doc,
        "The residual is divided equally among the four zones and held constant for every candidate. It is a bookkeeping term chosen so that the equal-split calibration case totals 15,400 lb O2/d. It does not model nitrification, endogenous respiration, alpha-factor variation or changing process conditions; therefore both the oxygen constraint and energy estimate remain screening calculations.",
        size=8.6,
    )
    add_heading(doc, "Machine-learning method and verification", level=2)
    body(
        doc,
        "Halton points are mapped to the four-part simplex by exponential normalization. Q, S0 and R span 3.5-8.0 MGD, 200-360 mg/L and 0.30-0.60. Inputs [f1,f2,f3,Q,S0,R] are standardized; f4 is dependent. Cubic polynomial ridge regression is selected after comparison with k-nearest neighbours and random Fourier features. The multi-output targets are soluble effluent BOD, four zone oxygen demands and total oxygen demand.",
        size=8.7,
    )
    add_equation(doc, "beta=(Phi^T Phi + 10^-6 I)^-1 Phi^T y;  Phi contains all standardized terms through degree 3")
    add_table(
        doc,
        ["Training labels", "Fixed test", "BOD RMSE", "Zone-O2 RMSE", "Decision stability"],
        [["4,096", "2,048 cases; 5 blocks", "0.0192 mg/L", "7.27 lb/d", "Same split as 8,192"]],
        [3.0, 4.2, 3.2, 3.8, 4.0],
        7.8,
    )
    body(
        doc,
        "At 8,192 labels the errors become 0.0191 mg/L and 7.21 lb/d: small non-monotonic fluctuations are expected across low-discrepancy blocks, while the recommendation is unchanged. Thus 4,096 is sufficient only for this surrogate, sampled range and mechanistic label generator. Low RMSE proves imitation of the assumed model, not accuracy at a real wastewater plant. The accompanying HU_TONGYU_MA_DINGNA_Code_Data.zip provides the scripts, sampled-result tables, units, standardization, fixed-test settings and reproduction steps.",
        size=8.5,
    )

    # Page 3: requested figure layout and reveal comparison.
    page(doc)
    add_heading(doc, "3  Design-space result and engineering reveal")
    pic = doc.add_picture(str(OUT / "capacity_value_curve.png"), width=Cm(17.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(
        doc,
        "Figure 1. A shows the screened BOD-oxygen design space and the diagnostic Zone-4-controlled band. B shows the minimum-soluble-BOD frontier as assumed local capacity C increases. C shows the corresponding four-zone oxygen demands and energy-only annual cost, calculated at 1.7 kg O2/kWh and $0.10/kWh.",
    )
    add_table(
        doc,
        ["Case", "Feed split f1/f2/f3/f4 (%)", "Modelled soluble BOD", "Peak / total O2 (lb/d)", "Engineering status"],
        [
            ["AI Round 1", "30.9/26.9/30.2/12.0", "6.69 mg/L", "4,426 / 15,598", "Aggressive; local capacity unknown"],
            ["C=4,200 candidate", "28.7/28.0/28.0/15.4", "7.79 mg/L", "4,200 / 15,547", "Pending field validation"],
            ["EPA reveal", "25/25/25/25", "11.00 mg/L", "3,850 / 15,400", "Calibration and plausibility benchmark"],
        ],
        [3.2, 4.5, 3.3, 3.8, 4.2],
        7.4,
    )
    body(
        doc,
        "The blue frontier answers a useful design question: how much predicted soluble-BOD benefit is unlocked as the maximum deliverable oxygen to any one zone, C, increases? Between C=3,850 and about 4,400 lb/d, front-loading becomes possible and predicted soluble BOD falls from 11.0 to about 6.64 mg/L. The flat section marks a change in the binding constraint: additional local distribution capacity no longer helps because the reported 15,600 lb/d total supply is nearly exhausted.",
        size=8.5,
    )
    body(
        doc,
        "C=4,200 is not 'conservative'. Its total demand of 15,547 lb/d leaves only 53 lb/d, or 0.34%, below the reported supply. It is therefore a candidate for field checking, not a design recommendation with a dependable safety margin. Load variability, oxygen-transfer efficiency and equipment ageing are not included.",
        lead="C=4,200 is not 'conservative'.",
        size=8.5,
    )

    # Page 4: evaluation and evidence boundary.
    page(doc)
    add_heading(doc, "4  Engineering evaluation after the reveal")
    body(
        doc,
        "The audit found that 3,900 lb O2/d per zone was never reported as hardware capacity; it came from dividing total supply by four. Imposing that value as four independent limits almost forces equal oxygen demand and biases the optimisation toward equal feed. The revised analysis treats C as unknown, retains total oxygen <=15,600 lb/d, and reports a capacity-response envelope instead of declaring one optimum.",
    )
    add_heading(doc, "What the reveal does and does not validate", level=2)
    body(
        doc,
        "EPA's 11 mg/L value was used to calibrate kappa and is then recovered by the equal-split model. This is internally consistent but circular as external validation. Moreover, the model state is soluble BOD, whereas the EPA value is carbonaceous BOD5 and the source separately reports TSS. The model omits particulate effluent BOD and final-clarifier performance. Therefore it can compare predicted effects of feed redistribution under fixed assumptions, but it cannot prove final-effluent compliance or superiority at the operating plant.",
        size=8.7,
    )
    add_table(
        doc,
        ["Finding", "Value", "Decision consequence"],
        [
            ["Local-capacity knee", "about 4,400 lb O2/d", "Beyond the knee, total supply binds"],
            ["C=4,200 oxygen margin", "53 lb/d = 0.34%", "Require variability and transfer tests"],
            ["Adverse case", "7.5 MGD; 320 mg/L; no feasible split", "Feed splitting cannot replace basin/blower capacity"],
            ["Energy-only change", "about $1.5k-$2.8k/y from C=3,850 to 4,400", "Not a lifecycle cost comparison"],
            ["Missing retrofit inputs", "diffusers, alpha, headloss, valves, blower curve", "Obtain field data and quotations"],
        ],
        [5.0, 4.8, 8.3],
        7.7,
    )
    body(
        doc,
        "The single cost line in Figure 1C uses 1.7 kg O2/kWh and $0.10/kWh. The reported $1.5k-$2.8k/y range is a separate sensitivity check using 2.2 and 1.2 kg O2/kWh, respectively. Both cover oxygen-transfer electricity only. Increasing C redistributes local delivery and does not automatically add total blower capacity. Complete annualized cost would require local diffuser/header/valve CAPEX, installation, fouling maintenance, pressure loss and blower operating-point data; therefore the figure is a sensitivity screen rather than an economic design.",
        size=8.6,
    )
    add_heading(doc, "Effectiveness and current engineering evidence", level=2)
    body(
        doc,
        "The useful AI outcome is not the 6.69 mg/L number alone. It is the discovery that local air distribution can be valuable until total oxygen binds, coupled with a falsifiable request for zone-level air-flow, DO, alpha-factor and hydraulic data. USIBWC reported operation of the newly constructed SAB plant in May 2025, while Minute 333 calls for further expansion study and real-time monitoring. Recent full-scale AI aeration studies also rely on sustained sensors, explicit baselines and plant trials. These developments support shadow-mode testing, but they do not transfer their performance gains to this step-feed case.",
        size=8.6,
    )

    # Page 5: integrated judgement, reflection and references.
    page(doc)
    add_heading(doc, "5  Overall assessment and professional reflection")
    body(
        doc,
        "Overall judgement. AI was effective as an auditable alternative generator and sensitivity-analysis tool. It searched 14,740 normal-mode splits, compared several surrogate methods, mechanistically recomputed shortlisted designs, and revealed the local-capacity frontier. Engineering review remained decisive: it removed an invented per-zone constraint, identified circular validation, distinguished soluble BOD from final cBOD5, and rejected an adverse case that no feed split could repair.",
        lead="Overall judgement.",
    )
    add_heading(doc, "Recommended next action", level=2)
    body(
        doc,
        "Do not implement C=4,200 as a final design. Use the 28.7/28.0/28.0/15.4% split as a commissioning hypothesis only after confirming zone air-delivery capacity. Install or verify zone air-flow and DO measurement, log valve position and header pressure, perform oxygen-transfer and alpha-factor testing, and run the candidate in shadow mode against the current equal split. Advance gradually only when effluent cBOD5/TSS, settling, DO, energy and peak-load performance remain acceptable with an explicit operating reserve.",
    )
    add_heading(doc, "Professional reflection", level=2)
    body(
        doc,
        "The most important lesson was that AI can make an unsupported assumption look precise. Dividing total oxygen by four appeared harmless, yet it created the main conclusion of the first analysis. Our responsibility is therefore to maintain an assumption register, check dimensions, separate calibration from validation, and define what evidence could disprove a recommendation. AI expands the option set and accelerates computation; the engineer remains accountable for model scope, constructability, safety margin, commissioning and regulatory compliance.",
        size=8.8,
    )
    body(
        doc,
        "In future work we would pre-register the objective and acceptance criteria, preserve code and test data, quantify uncertainty from influent and oxygen-transfer variation, and compare predictions with a sustained plant baseline. Human approval, manual override, rollback and drift monitoring would be mandatory. This turns AI from an answer generator into a transparent decision-support instrument.",
        size=8.8,
    )
    add_heading(doc, "References", level=2)
    compact_note(doc, "[1] U.S. EPA (2021), Project 8: Upgrade San Antonio de los Buenos WWTP, Addendum to the Feasibility Analysis. https://www.epa.gov/system/files/documents/2021-12/usmca-feasibility-analysis_project-8_upgrade-sab-wwtp-addendum.pdf")
    compact_note(doc, "[2] Oklahoma Administrative Code 252:656, Appendix A and Section 16-1, step-aeration screening ranges and oxygen criterion. https://www.law.cornell.edu/regulations/oklahoma/OAC-252-656-app-A")
    compact_note(doc, "[3] USIBWC (2025), Moving Fast and Fixing Things: infrastructure improvements and Tijuana River pollution. https://www.ibwc.gov/wp-content/uploads/2025/06/MOVING-FAST-AND-FIXING-THINGS-USIBWC-FOCUSED-ON-INFRASTRUCTURE-IMPROVEMENTS-AND-TIJUANA-RIVER-POLLUTION.pdf")
    compact_note(doc, "[4] U.S. EPA, Minute No. 333: expansion feasibility and real-time binational monitoring. https://www.epa.gov/sustainable-water-infrastructure/minute-no-333")
    compact_note(doc, "[5] Heo et al. (2025), real-scale digital-twin aeration optimisation, Desalination 593, 118235. https://doi.org/10.1016/j.desal.2024.118235")
    compact_note(doc, "[6] Xiong et al. (2025), ML-assisted expert aeration control, ACS ES&T Engineering 5, 3140-3152. https://doi.org/10.1021/acsestengg.5c00493")
    body(doc, f"Project repository: {GITHUB_URL}", lead="Project repository:", size=8.4, after=0)

    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
